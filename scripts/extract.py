"""One-shot extraction of May_2026_Meal_Plan.xlsx to TypeScript data modules.

Usage:
    python3 scripts/extract.py <path-to-xlsx>

Writes to ../data/*.ts relative to this script.
"""

from __future__ import annotations

import json
import re
import sys
from datetime import date
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = ROOT / "data"


def slugify(name: str) -> str:
    s = name.lower().strip()
    s = re.sub(r"[^a-z0-9]+", "-", s)
    return s.strip("-")


def expand_abbrev(text: str | None) -> str | None:
    if text is None:
        return None
    s = str(text)
    s = re.sub(r"\bIP\b", "Instant Pot", s)
    s = re.sub(r"\bOmega\s*3\b", "Omega-3", s)
    return s


def parse_int(v) -> int | None:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return int(v)
    m = re.search(r"\d+", str(v))
    return int(m.group()) if m else None


def parse_protein(v) -> int | None:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return int(v)
    s = str(v).strip()
    if s.lower() == "varies":
        return None
    m = re.search(r"\d+", s)
    return int(m.group()) if m else None


KNOWN_TOOLS = ["Instant Pot", "Air Fryer", "Sheet Pan", "Oven", "Sauté",
               "Stovetop", "Slow Cooker", "Grill"]


def extract_tools(tool: str | None) -> list[str]:
    if not tool:
        return []
    txt = expand_abbrev(tool) or ""
    parts = [p.strip() for p in re.split(r"\s*\+\s*", txt)]
    parts = [p for p in parts if not re.match(r"(?i)^(reuses|uses)\b", p)]
    found = []
    for part in parts:
        for t in KNOWN_TOOLS:
            if re.search(rf"\b{re.escape(t)}\b", part, re.IGNORECASE):
                if t not in found:
                    found.append(t)
    return found


def derive_tags(*, name: str, tool: str | None, base_tag: str | None,
                protein: int | None, prep_min: int | None, cook_min: int | None,
                ingredients: list[str], notes: str | None, category: str) -> list[str]:
    tags: list[str] = []

    for t in extract_tools(tool):
        if t not in tags:
            tags.append(t)

    if base_tag:
        parts = [p.strip() for p in re.split(r"\+", base_tag)]
        for p in parts:
            p = expand_abbrev(p) or ""
            p = p.strip()
            if p and p not in tags:
                tags.append(p)

    if protein is not None and protein >= 30:
        tags.append("High Protein")

    total = (prep_min or 0) + (cook_min or 0)
    if total and total <= 30:
        tags.append("Quick")

    ing_blob = " ".join(ingredients).lower() + " " + (notes or "").lower()
    name_blob = (name or "").lower()

    cuisine_map = [
        ("Mexican", ["taco", "carnita", "fajita", "enchilada", "barbacoa", "tortilla"]),
        ("Italian", ["parmesan", "parmigiana", "meatball", "wedding soup", "italian", "marinara"]),
        ("Mediterranean", ["mediterranean", "kalamata", "tzatziki", "hummus"]),
        ("Greek", ["greek chicken", "greek salad"]),
        ("BBQ", ["bbq", "pulled pork"]),
        ("Asian", ["soy sauce", "teriyaki"]),
    ]
    for label, kws in cuisine_map:
        if label in tags:
            continue
        if any(k in name_blob for k in kws) or any(k in ing_blob for k in kws if k not in ("greek",)):
            tags.append(label)
            break

    blob = ing_blob + " " + name_blob
    if any(k in blob for k in ["salmon", "sardine", "mackerel", "chia", "walnut"]):
        tags.append("Omega-3")
    if any(k in blob for k in ["lentil", "beef", "barley", "spinach", "apricot", "molasses"]):
        tags.append("Iron Rich")
    if category in ("breakfast",) and protein and protein >= 20:
        tags.append("Make Ahead")
    if any(k in blob for k in ["freeze", "freezer"]):
        if "Freezer Friendly" not in tags:
            tags.append("Freezer Friendly")

    if category == "soup":
        tags.append("Soup")
    if category == "snack":
        tags.append("Snack")
    if category == "lunch":
        tags.append("Lunch")
    if category == "breakfast":
        tags.append("Breakfast")
    if category == "dinner":
        tags.append("Dinner")

    seen = set()
    out = []
    for t in tags:
        if t and t not in seen:
            seen.add(t)
            out.append(t)
    return out


def parse_recipe_sheet(ws, category: str) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    recipes: list[dict] = []

    i = 2
    while i < len(rows):
        row = rows[i]
        a, b = (row[0], row[1]) if len(row) >= 2 else (None, None)
        if a and not b and not str(a).startswith(("Tool", "Protein", "Servings",
                                                   "Prep", "Cook", "Tag",
                                                   "Ingredients", "Method", "Step")):
            name = str(a).strip()
            recipe = {"name": name, "category": category, "ingredients": [], "steps": []}
            i += 1
            while i < len(rows):
                r = rows[i]
                a2 = r[0]
                b2 = r[1] if len(r) > 1 else None
                if a2 is None and b2 is None:
                    i += 1
                    break
                if a2 and not b2 and not str(a2).startswith(("Tool", "Protein",
                                                             "Servings", "Prep",
                                                             "Cook", "Tag",
                                                             "Ingredients", "Method", "Step")):
                    break
                key = str(a2).strip() if a2 else None
                val = b2
                if key == "Tool":
                    recipe["tool"] = expand_abbrev(val)
                elif key == "Protein per serving":
                    recipe["proteinGrams"] = parse_protein(val) or 0
                elif key == "Servings":
                    recipe["servings"] = parse_int(val)
                elif key == "Prep time":
                    recipe["prepMin"] = parse_int(val)
                elif key == "Cook time":
                    recipe["cookMin"] = parse_int(val)
                elif key == "Tag":
                    recipe["_baseTag"] = expand_abbrev(val)
                elif key == "Ingredients" and val is None:
                    pass
                elif key == "Method" and val is None:
                    pass
                elif a2 is None and isinstance(val, str):
                    cleaned = val.lstrip("•").strip()
                    if cleaned:
                        recipe["ingredients"].append(expand_abbrev(cleaned))
                elif key and key.startswith("Step") and isinstance(val, str):
                    recipe["steps"].append(expand_abbrev(val))
                i += 1
            recipes.append(recipe)
        else:
            i += 1
    return recipes


def parse_breakfasts(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    out = []
    for r in rows[3:]:
        name = r[0]
        if not name or str(name).startswith("Name"):
            continue
        protein = parse_protein(r[1])
        approach = r[2]
        ingredients_raw = r[3] or ""
        method = r[4] or ""
        ingredients = [expand_abbrev(x.strip()) for x in re.split(r",\s*", str(ingredients_raw)) if x.strip()]
        steps = [expand_abbrev(s.strip()) for s in re.split(r"(?<=\.)\s+", str(method)) if s.strip()]
        out.append({
            "name": str(name).strip(),
            "category": "breakfast",
            "proteinGrams": protein or 0,
            "notes": expand_abbrev(approach),
            "ingredients": ingredients,
            "steps": steps,
        })
    return out


def parse_lunches(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    out = []
    in_table = False
    for r in rows:
        a = r[0]
        if a == "Name":
            in_table = True
            continue
        if not in_table:
            continue
        if not a:
            continue
        protein = parse_protein(r[1])
        notes = r[2]
        out.append({
            "name": str(a).strip(),
            "category": "lunch",
            "proteinGrams": protein or 0,
            "notes": expand_abbrev(notes),
            "ingredients": [],
            "steps": [],
        })
    return out


def parse_snacks(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    out = []
    in_table = False
    blank_streak = 0
    for r in rows:
        a = r[0]
        if a == "Snack":
            in_table = True
            continue
        if not in_table:
            continue
        if not a:
            blank_streak += 1
            if blank_streak >= 1 and out:
                break
            continue
        if str(a).lower().startswith("avg "):
            break
        protein = parse_protein(r[1])
        notes = r[2]
        out.append({
            "name": str(a).strip(),
            "category": "snack",
            "proteinGrams": protein or 0,
            "notes": expand_abbrev(notes),
            "ingredients": [],
            "steps": [],
        })
    return out


def parse_calendar(ws, slug_index: dict[str, str]) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    out = []
    current_week_index = 0
    for r in rows:
        a = r[0]
        if not a:
            continue
        s = str(a).strip()
        m_week = re.match(r"Week\s+(\d+)", s)
        if m_week:
            current_week_index = int(m_week.group(1))
            continue
        if s in ("Date",) or "Daily Meal Calendar" in s:
            continue
        m = re.match(r"(\w+)\s+(\d+)", s)
        if not m:
            continue
        month_name, day_num = m.group(1), int(m.group(2))
        months = {"Jan": 1, "Feb": 2, "Mar": 3, "Apr": 4, "May": 5, "Jun": 6,
                  "Jul": 7, "Aug": 8, "Sep": 9, "Oct": 10, "Nov": 11, "Dec": 12}
        mo = months.get(month_name[:3])
        if not mo:
            continue
        d = date(2026, mo, day_num)
        iso_year, iso_week, _ = d.isocalendar()
        day_name = r[1] or d.strftime("%a")
        breakfast = expand_abbrev(r[2]) or ""
        lunch = expand_abbrev(r[3]) or ""
        dinner_name = str(r[4]).strip() if r[4] else ""
        notes = expand_abbrev(r[5]) or ""
        dinner_slug = slug_index.get(dinner_name, slugify(dinner_name))
        is_freezer_backup = bool(re.search(r"(backup|freeze)", notes, re.IGNORECASE))
        out.append({
            "date": d.isoformat(),
            "dayName": str(day_name).strip(),
            "week": iso_week,
            "weekIndex": current_week_index,
            "breakfast": breakfast,
            "lunch": lunch,
            "dinner": dinner_name,
            "dinnerSlug": dinner_slug,
            "notes": notes,
            "isFreezerBackup": is_freezer_backup,
        })
    return out


def parse_shopping(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    weeks: list[dict] = []
    current: dict | None = None
    current_section: str | None = None

    def start_week(label: str):
        nonlocal current, current_section
        m = re.match(r"Week\s+(\d+):\s*(.*)", label)
        num = int(m.group(1)) if m else len(weeks) + 1
        date_label = m.group(2).strip() if m else ""
        current = {"week": num, "dateLabel": date_label, "sections": []}
        current_section = None
        weeks.append(current)

    for r in rows:
        a = r[0]
        if not a:
            continue
        s = str(a).strip()
        if s.startswith("Week "):
            start_week(s)
            continue
        if s == "Item" or "Shopping Lists" in s or "Quantities are rough" in s:
            continue
        if current is None:
            continue
        is_section = (
            r[1] is None and r[2] is None and
            (s.isupper() or re.match(r"^[A-Z/ &\-]+$", s))
        )
        if is_section:
            current_section = s.title()
            current["sections"].append({"name": current_section, "items": []})
            continue
        if not current["sections"]:
            current["sections"].append({"name": "Misc", "items": []})
        item_name = expand_abbrev(s)
        qty = r[1]
        qty_str = str(qty).strip() if qty is not None else ""
        for_notes = expand_abbrev(r[2]) if r[2] else ""
        current["sections"][-1]["items"].append({
            "name": item_name,
            "qty": qty_str,
            "for": for_notes,
        })
    return weeks


def parse_prep(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    weeks: list[dict] = []
    current: dict | None = None
    for r in rows:
        a = r[0]
        if not a:
            continue
        s = str(a).strip()
        m = re.match(r"Week\s+(\d+)(?:\s*\((.*)\))?", s)
        if m:
            num = int(m.group(1))
            note = (m.group(2) or "").strip()
            current = {"week": num, "dateLabel": note, "steps": []}
            weeks.append(current)
            continue
        if s == "Step" or not current:
            continue
        if s.startswith("Step"):
            action = expand_abbrev(r[1]) if r[1] else ""
            approx = expand_abbrev(r[2]) if r[2] else ""
            current["steps"].append({
                "label": s,
                "action": action,
                "approx": approx or None,
            })
    return weeks


def parse_nutrition(ws) -> dict:
    rows = list(ws.iter_rows(values_only=True))
    intro = ""
    sections: list[dict] = []
    current: dict | None = None
    for i, r in enumerate(rows):
        a, b = r[0], r[1] if len(r) > 1 else None
        if a is None and b is None:
            continue
        s = str(a).strip() if a else ""
        if i == 0:
            continue
        if i == 1 and a and not b:
            intro = expand_abbrev(s)
            continue
        if a and not b:
            current = {"title": s, "items": []}
            sections.append(current)
            continue
        if a and b and current is not None:
            current["items"].append({
                "label": s,
                "value": expand_abbrev(str(b)),
            })
    return {"intro": intro, "sections": sections}


def ts_literal(obj) -> str:
    return json.dumps(obj, indent=2, ensure_ascii=False)


def main(xlsx_path: str) -> None:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)

    dinners = parse_recipe_sheet(wb["Dinner Recipes"], "dinner")
    soups = parse_recipe_sheet(wb["Soups"], "soup")
    breakfasts = parse_breakfasts(wb["Breakfasts"])
    lunches = parse_lunches(wb["Lunches & Wraps"])
    snacks = parse_snacks(wb["Snacks"])

    all_recipes = dinners + soups + breakfasts + lunches + snacks

    slug_index: dict[str, str] = {}
    out_recipes: list[dict] = []
    for r in all_recipes:
        s = slugify(r["name"])
        slug_index[r["name"]] = s
        tags = derive_tags(
            name=r["name"],
            tool=r.get("tool"),
            base_tag=r.pop("_baseTag", None),
            protein=r.get("proteinGrams"),
            prep_min=r.get("prepMin"),
            cook_min=r.get("cookMin"),
            ingredients=r.get("ingredients") or [],
            notes=r.get("notes"),
            category=r["category"],
        )
        entry = {
            "slug": s,
            "name": r["name"],
            "category": r["category"],
            "tool": r.get("tool"),
            "proteinGrams": r.get("proteinGrams") or 0,
            "servings": r.get("servings"),
            "prepMin": r.get("prepMin"),
            "cookMin": r.get("cookMin"),
            "tags": tags,
            "ingredients": r.get("ingredients") or [],
            "steps": r.get("steps") or [],
            "notes": r.get("notes"),
        }
        out_recipes.append(entry)

    calendar = parse_calendar(wb["May Calendar"], slug_index)
    shopping = parse_shopping(wb["Shopping Lists"])
    prep = parse_prep(wb["Sunday Prep Guide"])
    nutrition = parse_nutrition(wb["Nutrition Reference"])

    DATA_DIR.mkdir(parents=True, exist_ok=True)

    header = "// GENERATED by scripts/extract.py. Do not edit by hand.\n"

    (DATA_DIR / "recipes.ts").write_text(
        header +
        'import type { Recipe } from "./types";\n\n' +
        f"export const recipes: Recipe[] = {ts_literal(out_recipes)};\n"
    )

    (DATA_DIR / "calendar.ts").write_text(
        header +
        'import type { CalendarDay } from "./types";\n\n' +
        f"export const calendar: CalendarDay[] = {ts_literal(calendar)};\n"
    )

    (DATA_DIR / "shopping.ts").write_text(
        header +
        'import type { ShoppingWeek } from "./types";\n\n' +
        f"export const shopping: ShoppingWeek[] = {ts_literal(shopping)};\n"
    )

    (DATA_DIR / "prep.ts").write_text(
        header +
        'import type { PrepWeek } from "./types";\n\n' +
        f"export const prep: PrepWeek[] = {ts_literal(prep)};\n"
    )

    (DATA_DIR / "nutrition.ts").write_text(
        header +
        'import type { Nutrition } from "./types";\n\n' +
        f"export const nutrition: Nutrition = {ts_literal(nutrition)};\n"
    )

    print(f"Wrote {len(out_recipes)} recipes, {len(calendar)} days, "
          f"{len(shopping)} shopping weeks, {len(prep)} prep weeks.")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1
         else str(Path.home() / "Downloads" / "May_2026_Meal_Plan.xlsx"))
